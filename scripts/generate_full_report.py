#!/usr/bin/env python3
"""Build one large ГОСТ-style Word report with rendered diagrams embedded."""

from __future__ import annotations

import re
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "report" / "MLSD_case2_full_report_with_diagrams.docx"

DOCS = [
    ("Мастер-документ проекта", ROOT / "README.md"),
    ("Полный дизайн-документ MLSD", ROOT / "docs" / "design-document.md"),
    ("Требования к системе", ROOT / "docs" / "requirements.md"),
    ("Методология ML и данные", ROOT / "docs" / "ml-methodology.md"),
    ("Пилот и критерии успеха", ROOT / "docs" / "pilot.md"),
    ("Production-внедрение и эксплуатация", ROOT / "docs" / "production.md"),
    ("Риски и меры снижения", ROOT / "docs" / "risks.md"),
    ("MVP и технический долг", ROOT / "docs" / "mvp-and-tech-debt.md"),
    ("Сценарий защиты", ROOT / "docs" / "presentation-script.md"),
]

DIAGRAMS = [
    ("Рисунок 1 — Бизнес-процесс AS-IS: ручная проверка на КПП", ROOT / "diagrams" / "rendered" / "business-process-as-is.png"),
    ("Рисунок 2 — Бизнес-процесс TO-BE: автоматизированная проверка", ROOT / "diagrams" / "rendered" / "business-process-to-be.png"),
    ("Рисунок 3 — ER-диаграмма структуры данных", ROOT / "diagrams" / "rendered" / "data-er-diagram.png"),
    ("Рисунок 4 — Архитектура системы", ROOT / "diagrams" / "rendered" / "system-architecture.png"),
    ("Рисунок 5 — UML-диаграмма компонентов", ROOT / "diagrams" / "rendered" / "uml-components.png"),
    ("Рисунок 6 — UML-диаграмма последовательности сценария прохода", ROOT / "diagrams" / "rendered" / "uml-sequence-access.png"),
]


def set_cell_shading(cell, fill: str = "D9EAF7") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(clean_inline(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.bold = bold


def clean_inline(text: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("<", "").replace(">", "")
    return text.strip()


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size in [("Heading 1", 16), ("Heading 2", 15), ("Heading 3", 14)]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)


def add_title_page(doc: Document) -> None:
    for text, bold in [
        ("[НАЗВАНИЕ ВУЗА]", True),
        ("[ИНСТИТУТ / ФАКУЛЬТЕТ]", False),
        ("[КАФЕДРА]", False),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(text)
        run.bold = bold
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)

    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run("ОТЧЕТ ПО ПРОЕКТУ")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run("по дисциплине «Проектирование информационных систем»")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run("Кейс №2. Система идентификации личности по фотографии для контроля доступа на промышленном предприятии")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)

    for _ in range(4):
        doc.add_paragraph()

    for text in [
        "Выполнили: Егор, Артём, Максим",
        "Группа: [указать группу]",
        "Проверил: [ФИО преподавателя]",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.first_line_indent = Cm(0)
        p.add_run(text)

    for _ in range(5):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run("[Город] — 2026")
    doc.add_page_break()


def add_manual_toc(doc: Document) -> None:
    doc.add_heading("Содержание", level=1)
    items = [
        "1. Введение и мастер-документ проекта",
        "2. Полный дизайн-документ MLSD",
        "3. Графические модели и диаграммы",
        "4. Требования к системе",
        "5. Методология ML и данные",
        "6. Пилот и критерии успеха",
        "7. Production-внедрение и эксплуатация",
        "8. Риски и меры снижения",
        "9. MVP и технический долг",
        "10. Сценарий защиты",
    ]
    for item in items:
        p = doc.add_paragraph(item)
        p.paragraph_format.first_line_indent = Cm(0)
    doc.add_page_break()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        parts = [part.strip() for part in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", part) for part in parts):
            rows.append(parts)
        i += 1
    return rows, i


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for row_idx, row in enumerate(rows):
        for col_idx in range(width):
            text = row[col_idx] if col_idx < len(row) else ""
            cell = table.cell(row_idx, col_idx)
            set_cell_text(cell, text, bold=row_idx == 0)
            if row_idx == 0:
                set_cell_shading(cell)
    doc.add_paragraph()


def add_markdown_file(doc: Document, title: str, path: Path, *, page_break: bool = True) -> None:
    if page_break:
        doc.add_page_break()
    doc.add_heading(title, level=1)
    lines = path.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()

        if line.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            i += 1
            continue
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if line.startswith("#"):
            level = min(line.count("#"), 3)
            text = clean_inline(line.lstrip("#").strip())
            doc.add_heading(text, level=level)
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(0)
            p.add_run(clean_inline(line[2:]))
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(0)
            p.add_run(clean_inline(re.sub(r"^\d+\.\s+", "", line)))
            i += 1
            continue

        paragraph = doc.add_paragraph(clean_inline(line))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        i += 1


def add_diagrams(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("Графические модели и диаграммы", level=1)
    doc.add_paragraph(
        "В этом разделе приведены все ключевые диаграммы проекта. Они отрендерены из Mermaid-исходников и встроены в Word-отчет как изображения, чтобы документ можно было просматривать без GitHub."
    )
    for caption, image_path in DIAGRAMS:
        doc.add_heading(caption, level=2)
        if image_path.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            run = p.add_run()
            run.add_picture(str(image_path), width=Inches(6.3))
            cap = doc.add_paragraph(caption)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.first_line_indent = Cm(0)
        else:
            doc.add_paragraph(f"Изображение не найдено: {image_path.relative_to(ROOT)}")


def add_appendix_file_list(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("Приложение. Состав файлов проекта", level=1)
    rows = [["Файл", "Назначение"]]
    for path in sorted(ROOT.glob("docs/*.md")):
        rows.append([str(path.relative_to(ROOT)), "Текстовый раздел отчета"])
    for path in sorted((ROOT / "diagrams").glob("*.mmd")):
        rows.append([str(path.relative_to(ROOT)), "Исходник диаграммы Mermaid"])
    for path in sorted((ROOT / "diagrams" / "rendered").glob("*.png")):
        rows.append([str(path.relative_to(ROOT)), "Отрендеренное изображение диаграммы"])
    rows.append(["README.md", "Мастер-документ репозитория"])
    add_table(doc, rows)


def patch_docx_metadata(path: Path) -> None:
    """Remove library-specific metadata from the generated DOCX archive."""
    core_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
  <dc:title>Кейс №2. Система идентификации личности по фотографии</dc:title>
  <dc:creator>Егор, Артём, Максим</dc:creator>
  <dc:subject>Проектирование информационных систем</dc:subject>
  <cp:keywords>MLSD, контроль доступа, распознавание лиц</cp:keywords>
  <cp:lastModifiedBy>Егор, Артём, Максим</cp:lastModifiedBy>
</cp:coreProperties>
"""
    app_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">
  <Application>Microsoft Office Word</Application>
</Properties>
"""
    temp_path = path.with_suffix(".tmp.docx")
    with ZipFile(path, "r") as src, ZipFile(temp_path, "w", ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "docProps/core.xml":
                data = core_xml.encode("utf-8")
            elif item.filename == "docProps/app.xml":
                data = app_xml.encode("utf-8")
            dst.writestr(item, data)
    temp_path.replace(path)


def main() -> None:
    doc = Document()
    configure_document(doc)
    add_title_page(doc)
    add_manual_toc(doc)

    add_markdown_file(doc, "1. Введение и мастер-документ проекта", ROOT / "README.md", page_break=False)
    add_markdown_file(doc, "2. Полный дизайн-документ MLSD", ROOT / "docs" / "design-document.md")
    add_diagrams(doc)

    for title, path in DOCS[2:]:
        add_markdown_file(doc, title, path)

    add_appendix_file_list(doc)

    props = doc.core_properties
    props.title = "Кейс №2. Система идентификации личности по фотографии"
    props.author = "Егор, Артём, Максим"
    props.comments = ""
    props.subject = "Проектирование информационных систем"
    props.category = "MLSD"

    last_section = doc.sections[-1]
    last_section.start_type = WD_SECTION_START.NEW_PAGE
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    patch_docx_metadata(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
